from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import LayoutPlan


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Microsoft YaHei", size=10, color="263442")
THIN_BORDER = Border(bottom=Side(style="thin", color="B8C6D3"))


def _mm(units: int) -> float:
    return units / 100.0


def _payload(plan: LayoutPlan) -> dict:
    utilization = plan.metrics.panel_area / plan.metrics.opened_area if plan.metrics.opened_area else 0
    summary_rows = [
        ["字段", "值"],
        ["批次名称", plan.batch_name],
        ["优化档位", plan.mode],
        ["算法说明", plan.heuristic_note],
        ["板件数量", len(plan.panels)],
        ["板材总数", len(plan.sheets)],
        ["新增标准板", plan.metrics.new_standard_sheets],
        ["使用余料", plan.metrics.remnants_used],
        ["总利用率", utilization],
        ["废料面积(m²)", plan.metrics.waste_area / 10_000_000_000],
        ["有效零件间距(mm)", _mm(plan.settings.clearance)],
        ["板边留量(mm)", _mm(plan.settings.edge_margin)],
        ["暂定警告", plan.provisional_notice or "无"],
        ["分片策略", plan.split_strategy],
        ["焊缝数量", plan.weld_seams],
        ["完整箱面数量", len({panel.parent_panel_id or panel.panel_id for panel in plan.panels})],
        ["切割分片数量", len(plan.panels)],
    ]
    panel_groups = Counter(
        (
            panel.order_id,
            panel.product_id,
            panel.name,
            panel.source,
            panel.width,
            panel.height,
            panel.material_key,
            panel.rotation_allowed,
        )
        for panel in plan.panels
    )
    panel_rows = [[
        "订单号", "产品号", "板件名称", "来源", "长度(mm)", "宽度(mm)",
        "数量", "材料", "厚度(mm)", "表面要求", "方向规则",
    ]]
    for key, quantity in sorted(panel_groups.items(), key=lambda item: tuple(str(value) for value in item[0])):
        order_id, product_id, name, source, width, height, material, rotation_allowed = key
        panel_rows.append([
            order_id,
            product_id,
            name,
            source,
            _mm(width),
            _mm(height),
            quantity,
            material.material,
            _mm(material.thickness),
            material.surface,
            "允许90°旋转" if rotation_allowed else "方向锁定",
        ])
    panel_by_id = {panel.panel_id: panel for panel in plan.panels}
    placement_rows = [[
        "板件编号", "订单号", "产品号", "板件名称", "来源",
        "完整箱面编号", "完整箱面长度(mm)", "完整箱面宽度(mm)",
        "分片序号", "分片总数", "拼缝组", "板材编号",
        "X(mm)", "Y(mm)", "排版宽(mm)", "排版高(mm)", "旋转角度",
        "材料", "厚度(mm)", "表面要求",
    ]]
    for placement in plan.placements:
        panel = panel_by_id[placement.panel_id]
        placement_rows.append([
            panel.panel_id,
            panel.order_id,
            panel.product_id,
            panel.name,
            panel.source,
            panel.parent_panel_id or panel.panel_id,
            _mm(panel.parent_width or panel.width),
            _mm(panel.parent_height or panel.height),
            panel.piece_index,
            panel.piece_count,
            panel.seam_group or panel.panel_id,
            placement.sheet_id,
            _mm(placement.x),
            _mm(placement.y),
            _mm(placement.width),
            _mm(placement.height),
            placement.rotation,
            panel.material_key.material,
            _mm(panel.material_key.thickness),
            panel.material_key.surface,
        ])
    placements_by_sheet: dict[str, list] = {}
    for placement in plan.placements:
        placements_by_sheet.setdefault(placement.sheet_id, []).append(placement)
    stock_rows = [[
        "记录类型", "板材/余料编号", "来源规格", "材料", "厚度(mm)",
        "X(mm)", "Y(mm)", "长度(mm)", "宽度(mm)", "利用率",
    ]]
    for stock in plan.sheets:
        panel_area = sum(
            placement.width * placement.height
            for placement in placements_by_sheet.get(stock.sheet_id, [])
        )
        stock_rows.append([
            "标准板" if stock.kind == "standard" else "投入余料",
            stock.sheet_id,
            stock.stock_id,
            stock.material_key.material,
            _mm(stock.material_key.thickness),
            0,
            0,
            _mm(stock.width),
            _mm(stock.height),
            panel_area / (stock.width * stock.height),
        ])
    sheet_by_id = {sheet.sheet_id: sheet for sheet in plan.sheets}
    for index, remnant in enumerate(plan.remnants, start=1):
        stock = sheet_by_id[remnant.sheet_id]
        stock_rows.append([
            "新余料",
            f"{remnant.sheet_id}-R{index:03d}",
            remnant.sheet_id,
            stock.material_key.material,
            _mm(stock.material_key.thickness),
            _mm(remnant.x),
            _mm(remnant.y),
            _mm(remnant.width),
            _mm(remnant.height),
            None,
        ])
    return {
        "sheets": [
            {"name": "批次汇总", "tableName": "BatchSummary", "rows": summary_rows},
            {"name": "板件清单", "tableName": "PanelList", "rows": panel_rows},
            {"name": "排料坐标", "tableName": "PlacementCoordinates", "rows": placement_rows},
            {"name": "板材与余料", "tableName": "StockAndRemnants", "rows": stock_rows},
        ]
    }


def _write_sheet(workbook: Workbook, descriptor: dict) -> None:
    sheet = workbook.create_sheet(descriptor["name"])
    rows = descriptor["rows"]
    for row in rows:
        sheet.append(row)

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 25
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")
            cell.border = THIN_BORDER

    end = f"{get_column_letter(sheet.max_column)}{sheet.max_row}"
    table = Table(displayName=descriptor["tableName"], ref=f"A1:{end}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.auto_filter.ref = f"A1:{end}"


def _set_widths(sheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _set_number_format(sheet, cell_range: str, number_format: str) -> None:
    for row in sheet[cell_range]:
        for cell in row:
            cell.number_format = number_format


def _format_summary(sheet) -> None:
    _set_widths(sheet, {"A": 24, "B": 58})
    sheet["B9"].number_format = "0.00%"
    _set_number_format(sheet, "B10:B12", "0.00")


def _format_panels(sheet) -> None:
    _set_widths(
        sheet,
        {
            "A": 11,
            "B": 13,
            "C": 13,
            "D": 12,
            "E": 13,
            "F": 13,
            "G": 9,
            "H": 11,
            "I": 13,
            "J": 13,
            "K": 16,
        },
    )
    _set_number_format(sheet, f"E2:F{sheet.max_row}", "0.00")
    _set_number_format(sheet, f"G2:G{sheet.max_row}", "#,##0")
    _set_number_format(sheet, f"I2:I{sheet.max_row}", "0.00")


def _format_placements(sheet) -> None:
    _set_widths(
        sheet,
        {
            "A": 31,
            "B": 11,
            "C": 13,
            "D": 13,
            "E": 12,
            "F": 31,
            "G": 16,
            "H": 16,
            "I": 11,
            "J": 11,
            "K": 31,
            "L": 34,
            "M": 13,
            "N": 13,
            "O": 14,
            "P": 14,
            "Q": 11,
            "R": 11,
            "S": 13,
            "T": 13,
        },
    )
    _set_number_format(sheet, f"G2:H{sheet.max_row}", "0.00")
    _set_number_format(sheet, f"I2:J{sheet.max_row}", "0")
    _set_number_format(sheet, f"M2:P{sheet.max_row}", "0.00")
    _set_number_format(sheet, f"Q2:Q{sheet.max_row}", "0")
    _set_number_format(sheet, f"S2:S{sheet.max_row}", "0.00")


def _format_stock(sheet) -> None:
    _set_widths(
        sheet,
        {
            "A": 13,
            "B": 34,
            "C": 32,
            "D": 11,
            "E": 13,
            "F": 13,
            "G": 13,
            "H": 14,
            "I": 14,
            "J": 12,
        },
    )
    _set_number_format(sheet, f"E2:I{sheet.max_row}", "0.00")
    _set_number_format(sheet, f"J2:J{sheet.max_row}", "0.00%")


def export_excel(plan: LayoutPlan, path: Path) -> Path:
    payload = _payload(plan)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for descriptor in payload["sheets"]:
        _write_sheet(workbook, descriptor)
    _format_summary(workbook["批次汇总"])
    _format_panels(workbook["板件清单"])
    _format_placements(workbook["排料坐标"])
    _format_stock(workbook["板材与余料"])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
