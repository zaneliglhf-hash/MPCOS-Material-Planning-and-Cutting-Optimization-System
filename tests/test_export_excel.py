from openpyxl import load_workbook

from cutting_layout.export_excel import export_excel


def test_workbook_has_required_sheets_and_exact_placement_rows(valid_plan, tmp_path):
    path = export_excel(valid_plan, tmp_path / "parts.xlsx")
    workbook = load_workbook(path, data_only=True)
    assert workbook.sheetnames == ["批次汇总", "板件清单", "排料坐标", "板材与余料"]
    rows = list(workbook["排料坐标"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(valid_plan.placements)
    assert {row[0] for row in rows} == {placement.panel_id for placement in valid_plan.placements}
    assert workbook["批次汇总"]["B2"].value == valid_plan.batch_name


def test_coordinate_rows_match_validated_plan(valid_plan, tmp_path):
    path = export_excel(valid_plan, tmp_path / "parts.xlsx")
    workbook = load_workbook(path, data_only=True)
    rows = list(workbook["排料坐标"].iter_rows(min_row=2, values_only=True))
    row_by_panel = {row[0]: row for row in rows}
    for placement in valid_plan.placements:
        row = row_by_panel[placement.panel_id]
        assert row[11] == placement.sheet_id
        assert row[12] == placement.x / 100
        assert row[13] == placement.y / 100
        assert row[16] == placement.rotation


def test_panel_source_is_traceable(valid_plan, tmp_path):
    path = export_excel(valid_plan, tmp_path / "parts.xlsx")
    workbook = load_workbook(path, data_only=True)
    headers = [cell.value for cell in workbook["板件清单"][1]]
    source_column = headers.index("来源") + 1
    assert workbook["板件清单"].cell(2, source_column).value == "generated"


def test_waste_area_is_converted_from_centi_millimetres_to_square_metres(valid_plan, tmp_path):
    path = export_excel(valid_plan, tmp_path / "parts.xlsx")
    workbook = load_workbook(path, data_only=True)
    assert workbook["批次汇总"]["B10"].value == valid_plan.metrics.waste_area / 10_000_000_000


def test_workbook_exposes_warning_parent_and_piece_sequence(split_valid_plan, tmp_path):
    path = export_excel(split_valid_plan, tmp_path / "parts.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    summary = list(workbook["批次汇总"].values)
    assert any("禁止直接下料" in str(cell) for row in summary for cell in row)
    headers = next(workbook["排料坐标"].values)
    assert "完整箱面编号" in headers
    assert "分片序号" in headers
    assert "拼缝组" in headers


def test_export_excel_is_pure_python(valid_plan, tmp_path, monkeypatch):
    import subprocess

    monkeypatch.delenv("CUTTING_LAYOUT_NODE", raising=False)
    monkeypatch.delenv("CUTTING_LAYOUT_EXCEL_PREVIEW_DIR", raising=False)

    def reject_subprocess(*args, **kwargs):
        raise AssertionError("Excel export must not launch a subprocess")

    monkeypatch.setattr(subprocess, "run", reject_subprocess)
    path = export_excel(valid_plan, tmp_path / "portable.xlsx")
    assert path.exists()


def test_workbook_has_tables_filters_and_number_formats(valid_plan, tmp_path):
    path = export_excel(valid_plan, tmp_path / "styled.xlsx")
    workbook = load_workbook(path)
    expected_tables = {
        "批次汇总": "BatchSummary",
        "板件清单": "PanelList",
        "排料坐标": "PlacementCoordinates",
        "板材与余料": "StockAndRemnants",
    }
    for sheet_name, table_name in expected_tables.items():
        sheet = workbook[sheet_name]
        assert sheet.freeze_panes == "A2"
        assert table_name in sheet.tables
        assert sheet.auto_filter.ref is not None
        assert sheet.sheet_view.showGridLines is False
    assert workbook["批次汇总"]["B9"].number_format == "0.00%"
    assert workbook["排料坐标"]["M2"].number_format == "0.00"
