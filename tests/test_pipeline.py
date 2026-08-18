import json
import os
import subprocess
import sys
from pathlib import Path

import ezdxf
from openpyxl import load_workbook
import pytest

from cutting_layout.formats import parse_formats
from cutting_layout.output_transaction import OutputTransaction
from cutting_layout.pipeline import run_pipeline


def test_mixed_batch_creates_consistent_deliverables(tmp_path):
    input_path = Path("examples/mixed_batch.json")
    bundle = run_pipeline(input_path, tmp_path / "out")
    assert bundle.dxf_path.exists()
    assert bundle.excel_path.exists()
    assert bundle.pdf_path.exists()
    assert bundle.overview_path.exists()
    result = json.loads(bundle.result_json.read_text(encoding="utf-8"))
    assert result["panel_count"] == result["placement_count"]
    assert len(result["sheets"]) >= 1
    assert 0 < result["metrics"]["waste_area_m2"] < result["metrics"]["opened_area_m2"]
    dxf_count = len(ezdxf.readfile(bundle.dxf_path).modelspace().query('LWPOLYLINE[layer=="PART"]'))
    workbook = load_workbook(bundle.excel_path, read_only=True, data_only=True)
    excel_count = sum(
        1 for _ in workbook["排料坐标"].iter_rows(min_row=2, values_only=True)
    )
    assert dxf_count == excel_count == result["placement_count"]


def test_two_thousand_rectangles_finish_with_complete_placement(tmp_path, capacity_job_path):
    bundle = run_pipeline(capacity_job_path, tmp_path / "capacity-out")
    result = json.loads(bundle.result_json.read_text(encoding="utf-8"))
    assert result["panel_count"] == 2000
    assert result["placement_count"] == 2000


def test_cli_returns_two_and_does_not_create_false_success_files(tmp_path):
    bad_input = tmp_path / "bad.json"
    bad_input.write_text('{"batch_name": "broken"}', encoding="utf-8")
    completed = subprocess.run(
        [sys.executable, "-m", "cutting_layout", "run", str(bad_input), "--output", str(tmp_path / "out")],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 2
    assert not (tmp_path / "out" / "layout-all.dxf").exists()
    assert "error" in completed.stderr


def test_cli_exports_only_requested_formats(tmp_path):
    output = tmp_path / "selected"
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "cutting_layout",
            "run",
            "examples/mixed_batch.json",
            "--output",
            str(output),
            "--formats",
            "JSON,dxf,json",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0
    assert (output / "result.json").exists()
    assert (output / "layout-all.dxf").exists()
    assert not (output / "parts.xlsx").exists()
    assert not (output / "cutting-report.pdf").exists()
    assert not (output / "previews").exists()

    response = json.loads(completed.stdout)
    assert set(response) == {"status", "result_json", "dxf"}


def test_pdf_only_output_does_not_publish_previews(tmp_path):
    output = tmp_path / "pdf-only"

    bundle = run_pipeline(
        Path("examples/mixed_batch.json"),
        output,
        frozenset({"pdf"}),
    )

    assert bundle.pdf_path == output / "cutting-report.pdf"
    assert bundle.pdf_path.exists()
    assert not (output / "previews").exists()
    assert [item.name for item in output.iterdir()] == ["cutting-report.pdf"]


def test_transaction_restores_existing_output_when_publish_swap_fails(tmp_path, monkeypatch):
    output = tmp_path / "out"
    output.mkdir()
    sentinel = output / "keep.txt"
    sentinel.write_text("original", encoding="utf-8")
    real_replace = os.replace

    with pytest.raises(OSError, match="forced publish failure"):
        with OutputTransaction(output) as transaction:
            (transaction.staging_dir / "replacement.txt").write_text(
                "new",
                encoding="utf-8",
            )

            def fail_staging_publish(source, destination):
                if Path(source) == transaction.staging_dir:
                    raise OSError("forced publish failure")
                return real_replace(source, destination)

            monkeypatch.setattr(
                "cutting_layout.output_transaction.os.replace",
                fail_staging_publish,
            )
            transaction.commit()

    assert sentinel.read_text(encoding="utf-8") == "original"
    assert [item.name for item in output.iterdir()] == ["keep.txt"]
    assert list(tmp_path.glob(".out-staging-*")) == []
    assert list(tmp_path.glob(".out-backup-*")) == []


def test_cli_rejects_unknown_format(tmp_path):
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "cutting_layout",
            "run",
            "examples/mixed_batch.json",
            "--output",
            str(tmp_path / "out"),
            "--formats",
            "json,dwg",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 2
    assert "dwg" in completed.stderr


def test_cli_rejects_empty_format_list(tmp_path):
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "cutting_layout",
            "run",
            "examples/mixed_batch.json",
            "--output",
            str(tmp_path / "out"),
            "--formats",
            " , ",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 2
    assert "at least one" in completed.stderr


def test_parse_formats_normalizes_case_and_removes_duplicates():
    assert parse_formats(" PNG,Json,png ") == frozenset({"png", "json"})


def test_cli_version_is_available():
    completed = subprocess.run(
        [sys.executable, "-m", "cutting_layout", "--version"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0
    assert completed.stdout.strip() == "cutting-layout 0.2.0"


def test_failed_export_preserves_existing_output(tmp_path, monkeypatch):
    output = tmp_path / "out"
    output.mkdir()
    sentinel = output / "keep.txt"
    sentinel.write_text("original", encoding="utf-8")

    def fail_export(*args, **kwargs):
        raise RuntimeError("forced failure")

    monkeypatch.setattr("cutting_layout.pipeline.export_excel", fail_export)
    with pytest.raises(Exception, match="forced failure"):
        run_pipeline(Path("examples/mixed_batch.json"), output)

    assert sentinel.read_text(encoding="utf-8") == "original"
    assert [item.name for item in output.iterdir()] == ["keep.txt"]
    assert list(tmp_path.glob(".out-staging-*")) == []
    assert list(tmp_path.glob(".out-backup-*")) == []


def test_pipeline_selects_a_named_minimum_seam_split_strategy(tmp_path, oversize_job_path):
    bundle = run_pipeline(oversize_job_path, tmp_path / "out")
    result = json.loads(bundle.result_json.read_text(encoding="utf-8"))
    assert result["split_strategy"] in {"balanced", "prefer_1250", "prefer_1500"}
    assert result["weld_seams"] == 1
    assert result["parent_face_count"] == 1
    assert result["panel_count"] == result["placement_count"] == 2


def test_minimal_public_sample_runs_all_exporters(tmp_path):
    bundle = run_pipeline(Path("examples/minimal.json"), tmp_path / "minimal")
    result = json.loads(bundle.result_json.read_text(encoding="utf-8"))
    assert result["batch_name"] == "fictional-minimal-demo"
    assert result["panel_count"] == result["placement_count"] == 1
    assert len(result["sheets"]) == 1
    assert len(list((bundle.output_dir / "previews").glob("sheet-*.png"))) == 1
